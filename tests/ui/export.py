"""Browser checks for the owner's data export on the test-data preview (never the real site).

    python3 tools/build.py && python3 tools/preview/build_preview.py /tmp/pv/index.html
    (cd /tmp/pv && python3 -m http.server 8790 &) ; python3 tests/ui/export.py http://127.0.0.1:8790/ /tmp/pv/shots
"""
import asyncio, sys, os, json, zipfile
from playwright.async_api import async_playwright

BASE = sys.argv[1] if len(sys.argv) > 1 else "http://127.0.0.1:8790/"
SHOT = (sys.argv[2] if len(sys.argv) > 2 else "/tmp/pv/shots").rstrip("/") + "/"
os.makedirs(SHOT, exist_ok=True)
results = []


def check(name, cond, info=""):
    results.append((bool(cond), name, info))
    print(("PASS " if cond else "FAIL ") + name + ("" if cond else "  -> " + str(info)[:300]))


async def as_role(pg, role, hash="coach", reset=False):
    await pg.goto(BASE + "#home")
    await pg.evaluate("""([r,reset])=>{localStorage.setItem('tpe-preview-role',r);if(reset)localStorage.removeItem('tpe-preview-db-v6');localStorage.setItem('tpe-beta-ok','1')}""", [role, reset])
    await pg.goto("about:blank")
    await pg.goto(BASE + "#" + hash)
    await pg.wait_for_timeout(1200)


async def download(pg, label):
    async with pg.expect_download(timeout=20000) as dl:
        await pg.get_by_role("button", name=label).click()
    d = await dl.value
    path = SHOT + d.suggested_filename
    await d.save_as(path)
    return path


async def main():
    async with async_playwright() as p:
        b = await p.chromium.launch()
        ctx = await b.new_context(viewport={"width": 1200, "height": 900}, accept_downloads=True)
        pg = await ctx.new_page()
        errs = []
        pg.on("pageerror", lambda e: errs.append(str(e)))

        # teachers and students never see the export
        await as_role(pg, "teacher", reset=True)
        check("teacher: no export box", await pg.get_by_text("Бүх өгөгдлийг татах").count() == 0)
        await as_role(pg, "student")
        check("student: no export box", await pg.get_by_text("Бүх өгөгдлийг татах").count() == 0)

        # owner: everything
        await as_role(pg, "admin")
        box = pg.locator(".box", has_text="Бүх өгөгдлийг татах")
        check("admin: export box shown", await box.count() == 1)
        check("admin: 7 sections, all ticked", await box.locator("input[type=checkbox]:checked").count() == 7)
        await box.scroll_into_view_if_needed()
        await box.screenshot(path=SHOT + "export-box.png")

        x = await download(pg, "⬇ Excel")
        check("excel: file name", x.endswith(".xlsx") and "tomujin-pe-data-" in x, x)
        z = zipfile.ZipFile(x)
        check("excel: zip is valid", z.testzip() is None)
        import openpyxl
        wb = openpyxl.load_workbook(x)
        names = wb.sheetnames
        check("excel: summary first", names[0] == "Тойм", names)
        for want in ["Хэрэглэгчид", "Бэлтгэлийн тэмдэглэл", "Дугуйлангийн суудал", "Дугуйлангийн бүртгэл", "Дугуйланд ирэхгүй",
                     "Хүлээлгийн жагсаалт", "Чөлөөний хүсэлт", "Даалгавар", "Санал хүсэлт", "Саналын тэмдэглэл", "Тохиргоо"]:
            check("excel: sheet " + want, want in names, names)
        ws = wb["Бэлтгэлийн тэмдэглэл"]
        head = [c.value for c in ws[1]]
        rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
        check("loading: 5 rows", len(rows) == 5, len(rows))
        check("loading: student name column", "Нэр" in head and any(r[head.index("Нэр")] == "Бат Тэмүүлэн" for r in rows), head)
        check("loading: rpe is a number", all(isinstance(r[head.index("rpe")], (int, float)) for r in rows))
        check("loading: time readable", all(isinstance(r[head.index("at")], str) and r[head.index("at")][:2] == "20" for r in rows), rows[0])
        ws = wb["Дугуйлангийн бүртгэл"]
        head = [c.value for c in ws[1]]
        rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
        check("seats: club name + student", "Дугуйлангийн нэр" in head and any(r[head.index("Дугуйлангийн нэр")].startswith("Пиклбол") and r[head.index("Имэйл")] == "nomin.b@tomujin.edu.mn" for r in rows), (head, rows[:2]))
        sm = [[c.value for c in r] for r in wb["Тойм"].iter_rows()]
        check("summary: no read errors", not any(r and r[0] == "Уншиж чадаагүй хэсэг" for r in sm), sm)

        j = await download(pg, "⬇ JSON (бүтэн нөөц)")
        data = json.load(open(j, encoding="utf-8"))
        docs = data["docs"]
        check("json: meta", data["meta"]["app"] == "Tomujin PE" and data["meta"]["count"] == len(docs) and data["meta"]["errors"] == [], data["meta"])
        for pth in ["users/s1", "users/s1/loading/l0", "offers/tn-b/seats/s4", "offers/tn-b/seats/s4/abs/2026-11-19", "offers/jj-a/wait/w0",
                    "feedback/f1", "fbnotes/f2", "config/roles", "config/absence", "assign/a1", "site/clubreg"]:
            check("json: has " + pth, pth in docs)
        check("json: times marked", "__time" in json.dumps(docs["offers/jj-a/wait/w0"]), docs.get("offers/jj-a/wait/w0"))
        check("json: nothing outside the site's collections", not any(k.startswith("mock") for k in docs), [k for k in docs if k.startswith("mock")])

        # only the chosen sections
        await as_role(pg, "admin")
        box = pg.locator(".box", has_text="Бүх өгөгдлийг татах")
        for lbl in ["Бэлтгэлийн тэмдэглэл, үнэлгээ", "Хичээлээс чөлөө", "Сайтын агуулга", "Санал хүсэлт", "Тохиргоо"]:
            await box.locator("label", has_text=lbl).locator("input").uncheck()
        j = await download(pg, "⬇ JSON (бүтэн нөөц)")
        docs = json.load(open(j, encoding="utf-8"))["docs"]
        kinds = sorted({"/".join(k.split("/")[::2]) for k in docs})
        check("subset: only people + clubs", set(kinds) <= {"users", "pending", "offers", "offers/seats", "offers/seats/abs", "offers/wait", "offers/log", "regs"}, kinds)
        check("subset: users and seats present", "users/s1" in docs and "offers/tn-b/seats/s4" in docs, kinds)
        check("subset: no feedback/settings", not any(k.startswith(("feedback", "config", "private", "users/s1/loading")) for k in docs))
        st = await box.locator(".status").inner_text()
        check("status says done", "татагдлаа" in st, st)

        # nothing ticked
        for c in await box.locator("input[type=checkbox]").all():
            await c.uncheck()
        await pg.get_by_role("button", name="⬇ Excel").click()
        await pg.wait_for_timeout(300)
        st = await box.locator(".status").inner_text()
        check("nothing ticked: asks to choose", "сонгоно уу" in st, st)

        # phone width: no sideways scroll
        await pg.set_viewport_size({"width": 390, "height": 844})
        await pg.wait_for_timeout(300)
        sw = await pg.evaluate("document.documentElement.scrollWidth")
        check("phone: no sideways scroll", sw <= 391, sw)
        await box.screenshot(path=SHOT + "export-box-phone.png")

        check("no page errors", not errs, errs)
        await b.close()
    bad = [r for r in results if not r[0]]
    print(f"\n{len(results) - len(bad)}/{len(results)} passed")
    sys.exit(1 if bad else 0)

asyncio.run(main())
